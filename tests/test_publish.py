import json

from openpyxl import Workbook, load_workbook

from edinet_risk import publish


def _xlsx(p):
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); wb.active.append(["x"]); wb.save(p)


def _fixture(out):
    _xlsx(out / "run_r1" / "risk_diff_transport7.xlsx")
    _xlsx(out / "monitor_m1" / "monitor.xlsx")
    (out / "monitor_m1" / "monitor.md").write_text("flag: late なし\n", encoding="utf-8")
    (out / "run_r1" / "reconcile.md").write_text("照合OK\n", encoding="utf-8")
    item = {"heading_raw": "(1) 為替", "body": "為替変動のリスク。"}
    (out / "run_r1" / "compare.json").write_text(json.dumps({"filers": [
        {"label": "マツダ", "edinetCode": "E02163", "cur_doc": "S1", "prev_doc": "S0",
         "cur": {"period_end": "2026-03-31", "chars": 9, "items": [item]},
         "prev": {"period_end": "2025-03-31", "chars": 9, "items": [item]}},
        {"label": "他社", "edinetCode": "E99999", "skip": "no pair"}]}, ensure_ascii=False), encoding="utf-8")
    (out / "resolve_x").mkdir()
    (out / "resolve_x" / "resolve.csv").write_text(
        "key,name,edinet_code,cur_doc_id,cur_period_end,prev_doc_id,prev_period_end\n"
        "7261,マツダ,E02163,S1,2026-03-31,S0,2025-03-31\n1301,極洋,E00012,S9,2026-03-31,S8,2025-03-31\n",
        encoding="utf-8")


def test_build_and_sync(tmp_path):
    out, work, dest = tmp_path / "out", tmp_path / "publish", tmp_path / "sp"
    _fixture(out)
    tg = tmp_path / "t.toml"
    tg.write_text('[[filer]]\nlabel="マツダ"\nedinetCode="E02163"\n', encoding="utf-8")
    assert publish.build(out, work, [tg], dry_run=True)["ok"]
    assert not work.exists()
    res = publish.build(out, work, [tg], today="2026-10-29")
    assert res["ok"] and res["companies"] == 1
    for n in publish.CANONICAL:
        assert (work / n).exists()
    ws = load_workbook(work / "社一覧.xlsx").active
    assert [c.value for c in ws[1]] == publish.LIST_COLS
    assert ws.cell(2, 1).value == "7261" and ws.cell(2, 8).value == "2026-10-29"
    md = (work / "社別" / "7261_マツダ.md").read_text(encoding="utf-8")
    assert "為替変動のリスク" in md
    assert (work / "templates" / "yoY_report_template.xlsx").exists()
    assert "照合OK" in (work / "週次要約.md").read_text(encoding="utf-8")
    assert publish.sync(work, dest, dry_run=True) and not dest.exists()
    publish.sync(work, dest)
    assert (dest / "risk_diff.xlsx").exists()


def test_missing_source_keeps_previous(tmp_path):
    out, work = tmp_path / "out", tmp_path / "publish"
    out.mkdir()
    res = publish.build(out, work, [])
    assert not res["ok"] and "risk_diff" in res["missing"] and not work.exists()
