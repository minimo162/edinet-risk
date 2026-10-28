"""M2 Step H: `monitor` (plan §8-3 metrics, Excel + CSV, --fail-on-threshold -> exit 2; M2-8)."""
import csv
import json

import pytest
from openpyxl import load_workbook

from edinet_risk.__main__ import main
from edinet_risk.config import ConfigError
from edinet_risk.monitor import metrics
from edinet_risk.store import db

KEY = "MONITORKEY-0123456789abcdef"
RUN = "resolve-T1"
MONTH = "2026-09"
# entity, sec_code4, cur doc, prev doc
ENTS = (("E00001", "1001", "D0000001", "D0000002"), ("E00002", "1002", "D0000003", "D0000004"),
        ("E00003", "1003", "D0000005", "D0000006"))


def _sel(conn, ent, key, role, doc, *, flags=(), reason=None, no_prev=None, basis_extra=None):
    rb = {"target": {"key": key, "edinet_code": ent, "source": f"population:{MONTH}"}, **(basis_extra or {})}
    conn.execute("INSERT INTO filing_selection (run_id, entity_id, role, doc_id, period_end, flags_json, "
                 "no_prev_reason, reason_basis_json) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                 (RUN, ent, role, doc, "2026-03-31" if role == "cur" else "2025-03-31", json.dumps(list(flags)),
                  no_prev, json.dumps({"reason": reason if role == "cur" else no_prev, "reason_basis": rb})))


def make_db(tmp_path, ents=ENTS):
    """A resolve run where every entity is matched: cur/prev fetched ok, extracted, segmented, matched."""
    conn = db.open_db(tmp_path / "t.sqlite")
    db.start_run(conn, RUN, "resolve", code_version="t", fy="2026-03", as_of="2026-09-30",
                 args={"scope": f"population:{MONTH}"})
    db.finish_run(conn, RUN, "ok")
    with db.transaction(conn):
        for ent, sec, cur, prev in ents:
            conn.execute("INSERT INTO population (as_of_month, sec_code4, name, edinet_code, match_method) "
                         "VALUES (?, ?, ?, ?, 'sec_code')", (MONTH, sec, f"会社{sec}", ent))
            for doc, pe in ((cur, "2026-03-31"), (prev, "2025-03-31")):
                db.upsert(conn, "documents", {"doc_id": doc, "edinet_code": ent, "entity_id": ent, "xbrl_flag": "1",
                                              "doc_type_code": "120", "ordinance_code": "010", "form_code": "030000",
                                              "period_end": pe, "withdrawal_status": "0"}, key=["doc_id"])
                conn.execute("INSERT INTO fetches (doc_id, type, status, attempts) VALUES (?, 1, 'ok', 1)", (doc,))
                conn.execute("INSERT INTO extracts (doc_id, extractor_ver, input_sha256, source, text, char_len, "
                             "img_count, hit_count) VALUES (?, 'extract-1', 'x', 'xbrl', 't', 10000, 0, 1)", (doc,))
                for i in range(10):
                    conn.execute("INSERT INTO items (doc_id, segmenter_ver, item_idx, input_sha256, level, body, "
                                 "start_off, end_off, segmentation) VALUES (?, 'seg-1', ?, 'x', 1, 'b', 0, 1, "
                                 "'normal')", (doc, i))
            for i in range(10):
                conn.execute("INSERT INTO matches (cur_doc, prev_doc, matcher_ver, pair_idx, input_sha256, label, "
                             "review_flag) VALUES (?, ?, 'match-1', ?, 'x', 'same', ?)", (cur, prev, i, int(i < 2)))
            _sel(conn, ent, sec, "cur", cur)
            _sel(conn, ent, sec, "prev", prev)
        conn.execute("INSERT INTO api_calls (endpoint, params_masked, http_status, at) VALUES "
                     "('documents', '{}', 200, '2026-10-28T10:00:00+09:00')")
    return conn


def run_monitor(tmp_path, *extra, thresholds=None):
    args = ["monitor", "--db", str(tmp_path / "t.sqlite"), "--out", str(tmp_path / "out"),
            "--thresholds", str(thresholds or "config/monitor.toml"), *extra]
    return main(args)


def out_dir(tmp_path):
    return max((tmp_path / "out").glob("monitor_*"), key=lambda p: p.stat().st_mtime_ns)   # newest run


def summary(tmp_path) -> dict[str, dict]:
    with open(out_dir(tmp_path) / "monitor_summary.csv", encoding="utf-8-sig") as f:
        return {r["metric"]: r for r in csv.DictReader(f)}


# ---------------------------------------------------------------- exit codes
def test_within_thresholds_exit_0_and_outputs(tmp_path, capsys):
    make_db(tmp_path).close()
    assert run_monitor(tmp_path, "--fail-on-threshold") == 0
    d = out_dir(tmp_path)
    for name in ("monitor.xlsx", "monitor_summary.csv", "monitor_details.csv", "monitor.json", "monitor.md"):
        assert (d / name).exists()
    s = summary(tmp_path)
    assert float(s["match_rate"]["value"]) == 1.0 and s["match_rate"]["breached"] == "0"
    assert all(r["breached"] == "0" for r in s.values())
    wb = load_workbook(d / "monitor.xlsx")
    assert wb.sheetnames == ["監視_概要", "監視_詳細"]
    assert "breached=0" in capsys.readouterr().out


def test_breach_exit_2_only_with_flag(tmp_path, capsys):
    conn = make_db(tmp_path)
    with db.transaction(conn):
        conn.execute("INSERT INTO api_calls (run_id, endpoint, params_masked, http_status, at) VALUES "
                     "(NULL, 'documents', '{}', 429, '2026-10-28T11:00:00+09:00')")
    conn.close()
    assert run_monitor(tmp_path) == 0                          # report only
    assert run_monitor(tmp_path, "--fail-on-threshold") == 2
    s = summary(tmp_path)
    assert s["api_429"]["value"] == "1" and s["api_429_5xx"]["breached"] == "1"
    assert "BREACH api_429_5xx=1" in capsys.readouterr().out
    # --since after the 429 -> within thresholds again
    assert run_monitor(tmp_path, "--fail-on-threshold", "--since", "2026-10-28T12:00") == 0


@pytest.mark.parametrize("change", ["unknown", "fetch_failed", "empty", "multi_hit", "match_rate"])
def test_each_default_threshold_breaches(tmp_path, change):
    ents = ENTS
    if change == "match_rate":   # 100 entities, 2 not fetched -> 98% < 99%
        ents = tuple((f"E{i:05d}", f"{1000 + i}", f"D{2 * i:07d}", f"D{2 * i + 1:07d}") for i in range(1, 101))
    conn = make_db(tmp_path, ents)
    with db.transaction(conn):
        if change == "unknown":
            conn.execute("UPDATE filing_selection SET no_prev_reason = 'unknown', doc_id = NULL, reason_basis_json = "
                         "json_set(reason_basis_json, '$.reason', 'unknown', '$.reason_basis.unknown_cause', "
                         "'doclist_not_scanned') WHERE entity_id = 'E00001' AND role = 'prev'")
        elif change == "fetch_failed":
            conn.execute("UPDATE fetches SET status = 'failed', error_kind = 'not_found', last_http_status = 404 "
                         "WHERE doc_id = 'D0000002'")
        elif change == "empty":
            conn.execute("UPDATE extracts SET char_len = 0, hit_count = 0 WHERE doc_id = 'D0000001'")
        elif change == "multi_hit":
            conn.execute("UPDATE extracts SET hit_count = 2 WHERE doc_id = 'D0000003'")
        else:
            conn.execute("DELETE FROM fetches WHERE doc_id IN ('D0000002', 'D0000004')")   # cur docs of 2 entities
    conn.close()
    assert run_monitor(tmp_path, "--fail-on-threshold") == 2
    s = summary(tmp_path)
    mid = {"empty": "empty_extract"}.get(change, change)
    assert s[mid]["breached"] == "1"
    if change == "match_rate":
        assert float(s["match_rate"]["value"]) == pytest.approx(0.98)
        with open(out_dir(tmp_path) / "monitor_details.csv", encoding="utf-8-sig") as f:
            un = [r for r in csv.DictReader(f) if r["category"] == "unmatched"]
        assert [(r["entity_id"], r["metric"]) for r in un] == [("E00001", "unmatched.fetch_pending"),
                                                               ("E00002", "unmatched.fetch_pending")]


def test_thresholds_from_config(tmp_path):
    conn = make_db(tmp_path)
    with db.transaction(conn):    # one form_out_of_scope company: listed only by default
        conn.execute("UPDATE filing_selection SET reason_basis_json = json_set(reason_basis_json, "
                     "'$.reason_basis.out_of_scope_docs', json('[{\"doc_id\":\"X1\",\"docTypeCode\":\"120\","
                     "\"ordinanceCode\":\"010\",\"formCode\":\"032000\"}]')) WHERE entity_id = 'E00002' AND role = 'cur'")
    conn.close()
    assert run_monitor(tmp_path, "--fail-on-threshold") == 0
    assert summary(tmp_path)["form_out_of_scope"]["value"] == "1"
    cfg = tmp_path / "m.toml"
    cfg.write_text('version = "t"\n[thresholds]\nform_out_of_scope = 0\n', encoding="utf-8")
    assert run_monitor(tmp_path, "--fail-on-threshold", thresholds=cfg) == 2
    cfg.write_text('version = "t"\n[thresholds]\nmatch_rate = 1.0\n[entity]\nreview_ratio_max = 0.1\n',
                   encoding="utf-8")
    assert run_monitor(tmp_path, "--fail-on-threshold", thresholds=cfg) == 0   # review ratio is list-only
    s = summary(tmp_path)
    assert s["review_ratio_high"]["value"] == "3" and s["review_ratio_high"]["breached"] == "0"


def test_config_errors(tmp_path):
    p = tmp_path / "m.toml"
    p.write_text("[thresholds]\nno_such_metric = 1\n", encoding="utf-8")
    with pytest.raises(ConfigError):
        metrics.load_config(p)
    p.write_text("[entity]\nitems_min = \"x\"\n", encoding="utf-8")
    with pytest.raises(ConfigError):
        metrics.load_config(p)
    c = metrics.load_config("config/monitor.toml")
    assert c.thresholds["match_rate"] == 0.99 and c.entity.items_max == 60 and c.version == "m2-initial"


# ---------------------------------------------------------------- no API key in the outputs
def test_outputs_have_no_api_key(tmp_path, capsys, monkeypatch):
    monkeypatch.setenv("EDINET_API_KEY", KEY)
    conn = make_db(tmp_path)
    with db.transaction(conn):   # simulate a key that leaked into stored text
        conn.execute("UPDATE fetches SET status = 'failed', error_kind = 'invalid', last_http_status = 200, "
                     f"last_error = 'GET https://x/documents/D0000001?type=1&Subscription-Key={KEY} key {KEY}' "
                     "WHERE doc_id = 'D0000001'")
        conn.execute("INSERT INTO api_calls (endpoint, params_masked, http_status, at) VALUES "
                     "('documents', '{\"type\":\"1\"}', 503, '2026-10-28T11:00:00+09:00')")
    conn.close()
    assert run_monitor(tmp_path, "--fail-on-threshold") == 2
    out = capsys.readouterr()
    assert KEY not in out.out + out.err
    d = out_dir(tmp_path)
    for p in d.iterdir():
        if p.suffix == ".xlsx":
            wb = load_workbook(p)
            text = "\n".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
        else:
            text = p.read_text(encoding="utf-8-sig")
        assert KEY not in text, p.name
    with open(d / "monitor_details.csv", encoding="utf-8-sig") as f:
        fe = [r for r in csv.DictReader(f) if r["metric"] == "fetch_failed"]
    assert fe and "Subscription-Key=***" in fe[0]["detail"]


# ---------------------------------------------------------------- lists / M1 / formCode / metrics table
def test_lists_and_m1_anomalies(tmp_path):
    conn = make_db(tmp_path)
    with db.transaction(conn):
        conn.execute("UPDATE filing_selection SET flags_json = '[\"prev_gap\"]', reason_basis_json = json_set("
                     "reason_basis_json, '$.pair.gap_days', 10) WHERE entity_id = 'E00001' AND role = 'prev'")
        conn.execute("UPDATE filing_selection SET flags_json = '[\"irregular\"]' WHERE entity_id = 'E00002' "
                     "AND role = 'cur'")
        conn.execute("UPDATE extracts SET char_len = 1500, img_count = 3 WHERE doc_id = 'D0000003'")  # -85%, image
        conn.execute("DELETE FROM items WHERE doc_id = 'D0000005' AND item_idx > 0")                 # 1 item
        conn.execute("UPDATE items SET segmentation = 'fallback' WHERE doc_id = 'D0000005'")
        conn.execute("UPDATE matches SET review_flag = 1 WHERE cur_doc = 'D0000001' AND pair_idx < 5")
    conn.close()
    assert run_monitor(tmp_path, "--fail-on-threshold") == 0     # all of these are list-only
    s = summary(tmp_path)
    for mid, v in (("flag.prev_gap", "1"), ("flag.irregular", "1"), ("chars_prev_change", "1"),
                   ("image_low_text", "1"), ("items_out_of_range", "1"), ("items_prev_change", "1"),
                   ("segment_fallback", "1"), ("review_ratio_high", "1")):
        assert s[mid]["value"] == v, mid
    with open(out_dir(tmp_path) / "monitor_details.csv", encoding="utf-8-sig") as f:
        det = list(csv.DictReader(f))
    gap = next(r for r in det if r["metric"] == "flag.prev_gap")
    assert (gap["entity_id"], gap["key"], gap["name"], gap["value"]) == ("E00001", "1001", "会社1001", "10")
    conn = db.connect(tmp_path / "t.sqlite")
    run = out_dir(tmp_path).name.removeprefix("monitor_")
    rows = {(r[0], r[1]): r[2] for r in conn.execute("SELECT metric, entity_id, breached FROM metrics "
                                                     "WHERE run_id = ?", (run,))}
    assert rows[("match_rate", "")] == 0 and rows[("chars_prev_change.cur", "E00002")] == 1
    assert ("flag.prev_gap.prev", "E00001") in rows


def test_reasons_and_unmatched(tmp_path):
    conn = make_db(tmp_path)
    with db.transaction(conn):
        conn.execute("UPDATE filing_selection SET doc_id = NULL, reason_basis_json = json_set(reason_basis_json, "
                     "'$.reason', 'late') WHERE entity_id = 'E00003' AND role = 'cur'")
        conn.execute("UPDATE filing_selection SET doc_id = NULL, no_prev_reason = 'new_listing', reason_basis_json = "
                     "json_set(reason_basis_json, '$.reason', 'new_listing') WHERE entity_id = 'E00002' "
                     "AND role = 'prev'")
    conn.close()
    assert run_monitor(tmp_path, "--fail-on-threshold") == 2       # 2/3 matched < 99%
    s = summary(tmp_path)
    assert s["cur_reason.late"]["value"] == "1" and s["no_prev_reason.new_listing"]["value"] == "1"
    assert s["unmatched"]["value"] == "1" and s["unknown"]["value"] == "0"
    with open(out_dir(tmp_path) / "monitor_details.csv", encoding="utf-8-sig") as f:
        det = {(r["category"], r["entity_id"]): r for r in csv.DictReader(f)}
    assert det[("unmatched", "E00003")]["metric"] == "unmatched.late"
    assert det[("no_prev", "E00002")]["metric"] == "no_prev_reason.new_listing"


def test_excluded_form_change_vs_previous_run(tmp_path):
    conn = make_db(tmp_path)
    with db.transaction(conn):
        for i in range(10):
            db.upsert(conn, "documents", {"doc_id": f"X{i:07d}", "doc_type_code": "120", "ordinance_code": "010",
                                          "form_code": "080000", "period_end": "2025-12-31",
                                          "withdrawal_status": "0"}, key=["doc_id"])
    conn.close()
    assert run_monitor(tmp_path, "--fail-on-threshold") == 0          # first run: no previous
    assert summary(tmp_path)["excluded_form.120/010/080000"]["value"] == "10"
    conn = db.connect(tmp_path / "t.sqlite")
    with db.transaction(conn):
        for i in range(10, 30):
            db.upsert(conn, "documents", {"doc_id": f"X{i:07d}", "doc_type_code": "120", "ordinance_code": "010",
                                          "form_code": "080000", "period_end": "2025-12-31",
                                          "withdrawal_status": "0"}, key=["doc_id"])
    conn.close()
    assert run_monitor(tmp_path, "--fail-on-threshold") == 2          # 10 -> 30
    s = summary(tmp_path)
    assert s["excluded_form.120/010/080000"]["breached"] == "1"
    assert "030000" not in "".join(k for k in s if k.startswith("excluded_form."))   # in-scope form not listed


def test_monitor_without_resolve_run(tmp_path):
    db.open_db(tmp_path / "t.sqlite").close()
    assert run_monitor(tmp_path) == 2
