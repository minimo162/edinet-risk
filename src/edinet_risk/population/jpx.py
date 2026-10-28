"""`population load`: JPX 東証上場銘柄一覧 (data_j.xlsx) x EDINET codes -> ``population`` (plan §3, §5, §6-5).

Source: https://www.jpx.co.jp/markets/statistics-equities/misc/01.html -> ``data_j.xlsx`` (one sheet; header
``日付,コード,銘柄名,市場・商品区分,33業種コード,33業種区分,17業種コード,…``; ``日付`` is ``YYYYMMDD`` of the month
end). ``コード`` is an int for numeric codes and a string for codes with letters (``130A``). The legacy ``.xls``
format is not readable with openpyxl; see docs/jpx-population.md for converting it.

Only ``市場・商品区分 == プライム（内国株式）`` rows enter the population. Rows with a 5-character コード are
preferred / 社債型種類株式 of an issuer whose common stock is listed under the 4-character code (e.g. ``25935``
伊藤園第１種優先株式 next to ``2593``); they are excluded (the population is per company) and reported. Join with ``edinet_codes`` (``codes sync``):
candidates are the codes whose ``sec_code5 == sec_code4 || '0'``.
  * exactly one candidate -> ``match_method='sec_code'``
  * several, exactly one listed (上場) -> that one, ``match_method='sec_code_listed'`` (candidates reported)
  * several, not decidable -> ``edinet_code`` NULL, ``match_method='ambiguous'`` (candidates reported)
  * none -> ``edinet_code`` NULL, ``match_method='unmatched'`` (コード不一致, §6-5)

Reloading the same ``as_of_month`` replaces its rows in one transaction (idempotent).
"""
from __future__ import annotations

import hashlib
import re
import sqlite3
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from ..store.db import transaction

JPX_PAGE_URL = "https://www.jpx.co.jp/markets/statistics-equities/misc/01.html"
PRIME_DOMESTIC = "プライム（内国株式）"

_AS_OF_MONTH_RE = re.compile(r"^(\d{4})-(0[1-9]|1[0-2])$")
_SEC_CODE4_RE = re.compile(r"^[0-9][0-9A-Z]{3}$")
_CLASS_SHARE_RE = re.compile(r"^[0-9][0-9A-Z]{3}[0-9]$")
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_REQUIRED = ["日付", "コード", "銘柄名", "33業種区分"]
_MARKET_COLS = ["市場・商品区分", "市場"]


class JpxError(Exception):
    """Unusable JPX list or arguments."""


@dataclass
class JpxList:
    list_date: str | None      # 日付 as YYYY-MM-DD (None if the column is blank)
    total_rows: int
    markets: dict[str, int]    # 市場・商品区分 (as written) -> rows
    rows: list[dict]           # prime domestic only: sec_code4, name, market, sector33
    class_shares: list[tuple[str, str]]  # excluded prime rows with a 5-character コード: (code, 銘柄名)


def _norm(s) -> str:
    return unicodedata.normalize("NFKC", str(s)).strip()


def check_as_of_month(value: str) -> str:
    if not isinstance(value, str) or not _AS_OF_MONTH_RE.match(value):
        raise JpxError(f"--as-of-month must be YYYY-MM (e.g. 2026-09), got {value!r}")
    return value


def _sec_code4(v) -> str:
    if isinstance(v, bool):
        raise JpxError(f"コード {v!r}")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = _norm(v).upper()
    if not _SEC_CODE4_RE.match(s):
        raise JpxError(f"コード {v!r}")
    return s


def _list_date(v) -> str | None:
    if v is None or _norm(v) in ("", "-"):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = _norm(v).replace("/", "").replace("-", "")
    if isinstance(v, float) and v.is_integer():
        s = str(int(v))
    if not re.fullmatch(r"\d{8}", s):
        raise JpxError(f"日付 {v!r}")
    try:
        return date(int(s[:4]), int(s[4:6]), int(s[6:])).isoformat()
    except ValueError:
        raise JpxError(f"日付 {v!r}") from None


# ---------------------------------------------------------------- parser
def parse_jpx_xlsx(path: Path) -> JpxList:
    from openpyxl import load_workbook

    path = Path(path)
    if not path.is_file():
        raise JpxError(f"--jpx: {path} not found")
    with path.open("rb") as f:
        head = f.read(8)
    if head == _OLE2_MAGIC:
        raise JpxError(f"{path}: legacy .xls (OLE2) is not supported; save it as .xlsx first "
                       "(docs/jpx-population.md)")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as e:
        raise JpxError(f"{path}: not a readable .xlsx ({e})") from None
    try:
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    hi = col = None
    for i, row in enumerate(rows[:5]):
        cells = [_norm(c) if c is not None else "" for c in row]
        market = next((m for m in _MARKET_COLS if m in cells), None)
        if market and all(r in cells for r in _REQUIRED):
            hi, col = i, {r: cells.index(r) for r in _REQUIRED} | {"市場": cells.index(market)}
            break
    if hi is None:
        raise JpxError(f"{path}: header with columns {_REQUIRED} + 市場・商品区分 not found in the first rows")

    prime = _norm(PRIME_DOMESTIC)
    markets: dict[str, int] = {}
    dates: set[str] = set()
    out, seen, bad, class_shares = [], set(), [], []
    total = 0
    width = max(col.values()) + 1
    for n, row in enumerate(rows[hi + 1:], start=hi + 2):
        if not any(c is not None and _norm(c) for c in row):
            continue
        total += 1
        row = tuple(row) + (None,) * (width - len(row))
        market = "" if row[col["市場"]] is None else str(row[col["市場"]]).strip()
        markets[market] = markets.get(market, 0) + 1
        try:
            d = _list_date(row[col["日付"]])
            if d:
                dates.add(d)
            if _norm(market) != prime:
                continue
            raw_code = row[col["コード"]]
            if raw_code is not None and _CLASS_SHARE_RE.match(_norm(raw_code).upper()):
                class_shares.append((_norm(raw_code).upper(), str(row[col["銘柄名"]])))
                continue
            code = _sec_code4(raw_code)
        except JpxError as e:
            bad.append(f"row {n}: {e}")
            continue
        if code in seen:
            bad.append(f"row {n}: duplicate コード {code}")
            continue
        seen.add(code)
        sector = row[col["33業種区分"]]
        out.append({"sec_code4": code, "name": None if row[col["銘柄名"]] is None else str(row[col["銘柄名"]]),
                    "market": market, "sector33": None if sector is None else str(sector)})
    if bad:
        raise JpxError(f"{path}: {len(bad)} bad row(s): " + "; ".join(bad[:5]))
    if len(dates) > 1:
        raise JpxError(f"{path}: several 日付 values {sorted(dates)}")
    if not out:
        raise JpxError(f"{path}: no {PRIME_DOMESTIC} rows (markets: {sorted(markets)})")
    return JpxList(next(iter(dates), None), total, markets, out, class_shares)


# ---------------------------------------------------------------- join + load
def match_codes(conn: sqlite3.Connection, sec_code4s: list[str]) -> tuple[dict[str, tuple[str | None, str]],
                                                                         dict[str, list[dict]]]:
    """sec_code4 -> (edinet_code, match_method); plus candidates for the multi-candidate codes."""
    by_sec: dict[str, list[dict]] = {}
    for r in conn.execute("SELECT edinet_code, sec_code5, name, listed FROM edinet_codes "
                          "WHERE sec_code5 IS NOT NULL ORDER BY edinet_code"):
        by_sec.setdefault(r["sec_code5"], []).append(dict(r))
    result, multi = {}, {}
    for code in sec_code4s:
        cands = by_sec.get(code + "0", [])
        if not cands:
            result[code] = (None, "unmatched")
            continue
        if len(cands) == 1:
            result[code] = (cands[0]["edinet_code"], "sec_code")
            continue
        multi[code] = cands
        listed = [c for c in cands if c["listed"] == 1]
        result[code] = (listed[0]["edinet_code"], "sec_code_listed") if len(listed) == 1 else (None, "ambiguous")
    return result, multi


def load(conn: sqlite3.Connection, jpx_path: Path, as_of_month: str) -> dict:
    """Parse the JPX list, join it with edinet_codes and replace ``population`` for ``as_of_month``."""
    as_of_month = check_as_of_month(as_of_month)
    jpx = parse_jpx_xlsx(jpx_path)
    if jpx.list_date and jpx.list_date[:7] != as_of_month:
        raise JpxError(f"{jpx_path}: 日付 {jpx.list_date} is not in --as-of-month {as_of_month}")
    n_codes = conn.execute("SELECT count(*) FROM edinet_codes").fetchone()[0]
    if n_codes == 0:
        raise JpxError("edinet_codes is empty; run `codes sync` first")

    matched, multi = match_codes(conn, [r["sec_code4"] for r in jpx.rows])
    unlisted = {r["edinet_code"] for r in conn.execute("SELECT edinet_code FROM edinet_codes WHERE listed IS NOT 1")}
    rows = [{"as_of_month": as_of_month, **r, "edinet_code": matched[r["sec_code4"]][0],
             "match_method": matched[r["sec_code4"]][1]} for r in jpx.rows]
    with transaction(conn):
        replaced = conn.execute("DELETE FROM population WHERE as_of_month = ?", (as_of_month,)).rowcount
        conn.executemany("INSERT INTO population (as_of_month, sec_code4, name, market, sector33, edinet_code, "
                         "match_method) VALUES (:as_of_month, :sec_code4, :name, :market, :sector33, :edinet_code, "
                         ":match_method)", rows)
    methods: dict[str, int] = {}
    for r in rows:
        methods[r["match_method"]] = methods.get(r["match_method"], 0) + 1
    data = Path(jpx_path).read_bytes()
    return {
        "as_of_month": as_of_month, "list_date": jpx.list_date, "path": str(jpx_path),
        "sha256": hashlib.sha256(data).hexdigest(), "file_rows": jpx.total_rows, "markets": jpx.markets,
        "population": len(rows), "excluded_class_shares": jpx.class_shares, "replaced": replaced, "match_method": dict(sorted(methods.items())),
        "unmatched": [(r["sec_code4"], r["name"]) for r in rows if r["match_method"] == "unmatched"],
        "multi_candidates": {k: [c["edinet_code"] for c in v] for k, v in sorted(multi.items())},
        "matched_unlisted": [(r["sec_code4"], r["edinet_code"]) for r in rows
                             if r["edinet_code"] is not None and r["edinet_code"] in unlisted],
    }
