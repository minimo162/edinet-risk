"""Company diff workbook (plan §3 report/excel.py, §4 M1-4/M1-7, L1).

Sheets: ``概要`` (one row per filer), one sheet per filer (left/right item pairs with the character diff:
previous-only text in blue on the left, current-only text in red with underline on the right, via
openpyxl ``CellRichText``), ``照合`` (report.reconcile) and ``出典`` (source and processing notice).

A cell holds at most ``CELL_MAX`` characters (Excel: 32,767). Longer pairs continue on the next rows,
cut at sentence-block boundaries (a single over-long block is cut inside). Scores, label and stage
of each pair are separate columns (M1-7).

``snapshot()`` turns a written workbook into a list of (sheet, cell, value, style runs) for the
idempotency / snapshot tests (§8-2; the xlsx binary is never compared).
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill

from . import reconcile

CELL_MAX = 32_767
BLUE, RED = "FF0000FF", "FFFF0000"
PREV_FONT = InlineFont(color=BLUE)
CUR_FONT = InlineFont(color=RED, u="single")
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
LABEL_JA = {"same": "一致", "modified": "修正", "retitled": "改題", "moved": "移動", "merged": "統合",
            "split": "分割", "added": "新設", "deleted": "削除"}
CHANGE_JA = {"none": "なし", "minor": "軽微", "changed": "あり"}
PAIR_HEADER = ("#", "ラベル", "段階", "S", "本文類似", "見出し類似", "要確認", "番号変更", "変更", "軽微の種類",
               "前期の項目", "前期の本文（青字＝前期のみ）", "当期の項目", "当期の本文（赤字下線＝当期のみ）", "続き")
SUMMARY_HEADER = ("会社", "シート", "当期docID", "当期の版", "前期docID", "前期の版", "フラグ", "当期項目", "前期項目",
                  *[LABEL_JA[k] for k in LABEL_JA], "要確認", "変更あり", "軽微", "変更なし", "文単位の変更",
                  "うちLLM対象", "全文差分HTML", "備考")
HEAD_FILL = PatternFill("solid", fgColor="FFDDEBF7")

# Placeholder wording (plan D8: final text at the end of M1; review §5-6 example, PDL1.0 / [規約] 1.1).
SOURCE_NOTICE = (
    ("出典", "EDINET閲覧（提出）サイト（https://disclosure2.edinet-fsa.go.jp/）。金融庁 EDINET API v2 により取得した"
             "有価証券報告書の「事業等のリスク」（jpcrp_cor:BusinessRisksTextBlock）"),
    ("利用条件", "公共データ利用規約（第1.0版）（PDL1.0）に基づき利用"),
    ("加工した旨", "本資料は上記データを加工して作成したもの（本文の抽出・項目分割・前年との対応付け・文字単位の差分の"
                "色付け）。表記の正規化は比較時のみに行い、本文は原文のまま表示している"),
    ("加工者", "（仮）作成部署名をここに記載する"),
    ("注意", "文言は計画 D8 の仮置き（M1の終わりに確定）。自動処理の結果であり、正確性は手作業の照合で確認すること"),
)


def clean(s: str) -> str:
    return _ILLEGAL.sub("", s)


def cell_len(s: str) -> int:
    """Length as Excel counts it (UTF-16 code units)."""
    return len(s.encode("utf-16-le")) // 2


def sheet_title(label: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", label)[:28] or "sheet"
    t, k = base, 2
    while t in used:
        t, k = f"{base[:26]}_{k}", k + 1
    used.add(t)
    return t


def _side_runs(block: dict, side: str) -> list[tuple[str, str]]:
    own = "prev_only" if side == "prev" else "cur_only"
    runs = [("equal", block[f"gap_{side}"])] if block[f"gap_{side}"] else []
    runs += [(op, t) for op, t in block["segments"] if op in ("equal", own) and t]
    return runs


def _split_runs(runs: list[tuple[str, str]], limit: int) -> list[list[tuple[str, str]]]:
    out, cur, n = [], [], 0
    for op, t in runs:
        while t:
            room = limit - n
            if room <= 0:
                out.append(cur)
                cur, n = [], 0
                continue
            piece = t[:room]
            while cell_len(piece) > room:           # astral characters count twice
                piece = piece[:-1]
            if not piece:                           # an astral char does not fit in the last unit
                out.append(cur)
                cur, n = [], 0
                continue
            cur.append((op, piece))
            n += cell_len(piece)
            t = t[len(piece):]
    if cur:
        out.append(cur)
    return out


def chunk_pair(blocks: list[dict], limit: int = CELL_MAX) -> list[tuple[list, list]]:
    """[(prev_runs, cur_runs)] per output row; each side of a row has at most ``limit`` characters."""
    rows: list[tuple[list, list]] = []
    prev, cur, np_, nc = [], [], 0, 0
    for b in blocks:
        pr, cr = _side_runs(b, "prev"), _side_runs(b, "cur")
        lp, lc = sum(cell_len(t) for _, t in pr), sum(cell_len(t) for _, t in cr)
        if (np_ + lp > limit or nc + lc > limit) and (prev or cur):
            rows.append((prev, cur))
            prev, cur, np_, nc = [], [], 0, 0
        if lp > limit or lc > limit:
            ps, cs = _split_runs(pr, limit), _split_runs(cr, limit)
            for k in range(max(len(ps), len(cs))):
                rows.append((ps[k] if k < len(ps) else [], cs[k] if k < len(cs) else []))
            continue
        prev += pr
        cur += cr
        np_, nc = np_ + lp, nc + lc
    if prev or cur or not rows:
        rows.append((prev, cur))
    return rows


def rich(runs: list[tuple[str, str]]) -> CellRichText | str | None:
    merged: list[list] = []
    for op, t in runs:
        t = clean(t)
        if not t:
            continue
        style = "equal" if op == "equal" else op
        if merged and merged[-1][0] == style:
            merged[-1][1] += t
        else:
            merged.append([style, t])
    if not merged:
        return None
    if all(s == "equal" for s, _ in merged):
        return "".join(t for _, t in merged)
    return CellRichText([t if s == "equal" else TextBlock(PREV_FONT if s == "prev_only" else CUR_FONT, t)
                         for s, t in merged])


def _header(ws, header, widths=None) -> None:
    ws.append(list(header))
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    for col, w in (widths or {}).items():
        ws.column_dimensions[col].width = w


def _heads(items: list[dict], idxs: list[int]) -> str:
    by = {it["idx"]: it["heading_raw"] for it in items}
    return " ＋ ".join(clean(by[i]) for i in idxs)


def _r(x):
    return None if x is None else round(x, 4)


def write_filer_sheet(ws, f: dict, cell_max: int = CELL_MAX) -> None:
    _header(ws, PAIR_HEADER, {"A": 5, "B": 7, "C": 8, "K": 24, "L": 70, "M": 24, "N": 70, "O": 8})
    top = Alignment(wrap_text=True, vertical="top")
    for p in f["pairs"]:
        chunks = chunk_pair(p["blocks"], cell_max)
        for k, (pr, cr) in enumerate(chunks):
            first = k == 0
            row = [p["pair_idx"] + 1,
                   LABEL_JA[p["label"]] if first else None, p["stage"] if first else None,
                   _r(p["score_s"]) if first else None, _r(p["body_sim"]) if first else None,
                   _r(p["head_sim"]) if first else None,
                   ("○" if p["review_flag"] else None) if first else None,
                   ("○" if p["renumbered"] else None) if first else None,
                   CHANGE_JA[p["change"]] if first else None, p["minor_kind"] if first else None,
                   (_heads(f["prev"]["items"], p["prev_items"]) or None) if first else None, rich(pr),
                   (_heads(f["cur"]["items"], p["cur_items"]) or None) if first else None, rich(cr),
                   f"{k + 1}/{len(chunks)}" if len(chunks) > 1 else None]
            ws.append(row)
            for c in ws[ws.max_row]:
                c.alignment = top


def summary_row(f: dict, sheet: str | None, html_name: str | None) -> list:
    pairs = f["pairs"]
    lab = {k: sum(p["label"] == k for p in pairs) for k in LABEL_JA}
    ch = {k: sum(p["change"] == k for p in pairs) for k in CHANGE_JA}
    ok = f["cur"] is not None
    return [f["label"], sheet, f["cur_doc"], f["cur_kind"], f["prev_doc"], f["prev_kind"], ", ".join(f["flags"]) or None,
            len(f["cur"]["items"]) if ok else None, len(f["prev"]["items"]) if ok else None,
            *[lab[k] if ok else None for k in LABEL_JA], sum(p["review_flag"] for p in pairs) if ok else None,
            ch["changed"] if ok else None, ch["minor"] if ok else None, ch["none"] if ok else None,
            sum(p["n_changes"] for p in pairs) if ok else None, sum(p["n_changes_llm"] for p in pairs) if ok else None,
            html_name, f["skip"]]


def write_workbook(path: Path, data: dict, *, html_names: dict[str, str] | None = None, gold: dict | None = None,
                   cell_max: int = CELL_MAX) -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "概要"
    _header(ws, SUMMARY_HEADER, {"A": 16, "B": 16, "G": 18, "AA": 24, "AB": 40})
    used = {"概要", "照合", "出典"}
    for f in data["filers"]:
        name = sheet_title(f["label"], used) if f["cur"] is not None else None
        ws.append(summary_row(f, name, (html_names or {}).get(f["label"])))
        if name:
            write_filer_sheet(wb.create_sheet(name), f, cell_max)
    rs = reconcile.rows(data, gold)
    wr = wb.create_sheet("照合")
    _header(wr, reconcile.HEADER, {"A": 16, "B": 12, "H": 36, "I": 36, "O": 40})
    for r in rs:
        wr.append(list(r))
    write_source_sheet(wb.create_sheet("出典"), data)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"reconcile": reconcile.summary(rs)}


def write_source_sheet(ws, data: dict) -> None:
    _header(ws, ("項目", "内容"), {"A": 18, "B": 110})
    for k, v in SOURCE_NOTICE:
        ws.append([k, v])
    ws.append([])
    ws.append(["凡例", "前期の本文の青字＝前期のみの文字（削除）、当期の本文の赤字下線＝当期のみの文字（追加）"])
    ws.append(["処理", f"run_id={data['run_id']} {data['compare_ver']} {data['segmenter_ver']} {data['matcher_ver']} "
                      f"{data['differ_ver']} {data['minor_ver']} thresholds={data['thresholds_version']}"
                      f"({data['thresholds_fingerprint']}) as_of={data['as_of']}"])
    for f in data["filers"]:
        docs = [d for d in (f["prev_doc"], f["cur_doc"]) if d]
        ws.append([f["label"], "docID " + " → ".join(docs) + (f"（{f['skip']}）" if f["skip"] else "")])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


# ---------------------------------------------------------------- snapshot (tests, §8-2)
def _font(f) -> tuple:
    if f is None:
        return (None, None, None)
    color = getattr(getattr(f, "color", None), "rgb", None)
    return (color if isinstance(color, str) else None, getattr(f, "u", None) or getattr(f, "underline", None),
            bool(getattr(f, "b", None)))


def snapshot(path: Path, mask: dict[str, str] | None = None) -> list:
    """[(sheet, coord, [(text, color, underline, bold)...]) for every non-empty cell], strings masked."""
    def m(s):
        if isinstance(s, str):
            for k, v in (mask or {}).items():
                s = s.replace(k, v)
        return s

    wb = load_workbook(path, rich_text=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, CellRichText):
                    runs = [(m(r), None, None, False) if isinstance(r, str) else (m(r.text), *_font(r.font))
                            for r in v]
                else:
                    runs = [(m(v), *_font(c.font))]
                out.append((ws.title, c.coordinate, runs))
    return out
